import csv
import datetime
import io
import json
import secrets
import tempfile
import time
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from adminsite.backends.sqlalchemy.values import to_column_type
from adminsite.exceptions import AdminSiteError, FieldValidationError
from adminsite.fields import ChoiceField, FileField, RelationField
from adminsite.i18n import gettext as _
from adminsite.views import ModelView

# Files waiting between the preview and the import are kept this long.
PLAN_LIFETIME = 60 * 60
PLAN_FOLDER = Path(tempfile.gettempdir()) / "adminsite-imports"
YES_OR_NO = frozenset({"", "1", "0", "true", "false", "yes", "no", "on", "off"})


class ImportProblem(AdminSiteError):
    """The file could not be read as a table."""


def read_table(filename: str, data: bytes) -> list[list[str]]:
    """Read a CSV or Excel file into rows of text, header row first."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        rows = _read_excel(data)
    elif suffix in (".csv", ".txt", ""):
        rows = _read_csv(data)
    else:
        raise ImportProblem(_("Choose a CSV file or an Excel file (.xlsx)."))
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        raise ImportProblem(_("The file is empty."))
    return rows


def _read_csv(data: bytes) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Windows saves CSV in the old Windows code page.
        text = data.decode("cp1252", errors="replace")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _read_excel(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportProblem(
            _("Reading Excel files needs openpyxl: pip install 'adminsite[excel]'.")
        ) from None
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportProblem(
            _("This file could not be read as an Excel file.")
        ) from None
    try:
        sheet = book.worksheets[0]
        return [
            [cell_text(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        book.close()


def cell_text(value: Any) -> str:
    """Write a spreadsheet cell the way a form would have sent it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime | datetime.date | datetime.time):
        return value.isoformat()
    return str(value)


@dataclass
class ImportRow:
    """One row of the file, checked against the view."""

    number: int
    raw: dict[str, str]
    key: str = ""
    errors: dict[str, str] = field(default_factory=dict)
    values: dict[str, Any] = field(default_factory=dict)

    @property
    def action(self) -> str:
        """What importing does with the row: create, update or nothing."""
        if self.errors:
            return "error"
        return "update" if self.key else "create"


@dataclass
class ImportPlan:
    """A file matched to a view's fields, every row checked."""

    columns: list[str]
    ignored: list[str]
    rows: list[ImportRow]

    def count(self, action: str) -> int:
        """How many rows would be created, updated, or have errors."""
        return sum(1 for row in self.rows if row.action == action)

    @property
    def ready(self) -> list[ImportRow]:
        """The rows that can be imported."""
        return [row for row in self.rows if not row.errors]


def import_columns(view: ModelView, request: Any = None) -> tuple[str, ...]:
    """The paths a file can fill: the key, then the form's own fields."""
    readonly = set(view.get_readonly_fields(request))
    fields = tuple(
        path
        for path in view.get_form_fields(request)
        if path not in readonly
        and view.field_for(path).stored
        and not isinstance(view.field_for(path), FileField)
        and not _is_collection(view, path)
    )
    key = view.schema.primary_key
    return (key[0], *fields) if len(key) == 1 and key[0] not in fields else fields


def _is_collection(view: ModelView, path: str) -> bool:
    item = view.field_for(path)
    return isinstance(item, RelationField) and item.collection


def match_headers(
    view: ModelView, headers: Sequence[str], request: Any = None
) -> tuple[list[str | None], list[str]]:
    """Match each header to a path, by its name or its label, ignoring case."""
    known: dict[str, str] = {}
    for path in import_columns(view, request):
        known[path.lower()] = path
        known[view.label_for(path).strip().lower()] = path
    matched: list[str | None] = []
    ignored = []
    for header in headers:
        found = known.get(header.strip().lower())
        if found is None or found in matched:
            matched.append(None)
            if header.strip():
                ignored.append(header.strip())
        else:
            matched.append(found)
    return matched, ignored


async def build_plan(
    view: ModelView,
    session: Any,
    table: Sequence[Sequence[str]],
    *,
    request: Any = None,
) -> ImportPlan:
    """Check every row of a table against the view, without writing anything."""
    header, *body = table
    matched, ignored = match_headers(view, header, request)
    columns = [path for path in matched if path is not None]
    if not columns:
        raise ImportProblem(
            _("None of the columns match. Use the field names from the template.")
        )

    rows = []
    for offset, cells in enumerate(body):
        raw = {
            path: (cells[index] if index < len(cells) else "").strip()
            for index, path in enumerate(matched)
            if path is not None
        }
        rows.append(ImportRow(number=offset + 2, raw=raw))

    key_name = view.schema.primary_key[0] if len(view.schema.primary_key) == 1 else ""
    existing = await _existing(view, session, rows, key_name, request)
    for row in rows:
        check_row(view, row, key_name, existing, request)
    return ImportPlan(columns, ignored, rows)


async def _existing(
    view: ModelView,
    session: Any,
    rows: Sequence[ImportRow],
    key_name: str,
    request: Any,
) -> dict[str, Any]:
    """The records the file names by key, found in one query per 500."""
    if not key_name:
        return {}
    column = getattr(view.model, key_name)
    python_type = view.schema.field_named(key_name).python_type
    wanted = []
    for key in {row.raw[key_name] for row in rows if row.raw.get(key_name)}:
        try:
            wanted.append(to_column_type(python_type, key))
        except ValueError:
            continue
    found: dict[str, Any] = {}
    scope = view.scope_for(request)
    for start in range(0, len(wanted), 500):
        statement = view.repository.base_statement(scope).where(
            column.in_(wanted[start : start + 500])
        )
        for record in (await session.scalars(statement)).all():
            found[str(getattr(record, key_name))] = record
    return found


def check_row(
    view: ModelView,
    row: ImportRow,
    key_name: str,
    existing: dict[str, Any],
    request: Any = None,
) -> None:
    """Read one row into values, noting what is wrong with it."""
    key = row.raw.get(key_name, "") if key_name else ""
    record = None
    if key:
        record = existing.get(key)
        if record is None:
            row.errors[key_name] = _(
                "No {thing} has the key {key}.", thing=view.label.lower(), key=key
            )
            return
        row.key = key

    for path, text in row.raw.items():
        if path == key_name:
            continue
        item = view.field_for(path)
        try:
            row.values[path] = item.parse(normalize(item, text))
        except FieldValidationError as error:
            row.errors[path] = error.message

    if record is None:
        for path in import_columns(view, request):
            if path == key_name or path in row.raw:
                continue
            if view.field_for(path).required:
                row.errors[path] = _("Missing, and a new record needs it.")


def normalize(item: Any, text: str) -> str:
    """Accept what an export or a spreadsheet writes: labels, grouped numbers.

    A form's checkbox reads anything but a yes as no, but a file that says
    "maybe" is a mistake worth pointing out.
    """
    if item.widget == "checkbox" and text.lower() not in YES_OR_NO:
        raise FieldValidationError(item.name, _("Write yes or no."))
    if isinstance(item, ChoiceField):
        for value, label in item.choices:
            if text.lower() == label.lower():
                return value
    if item.widget == "number":
        return text.replace(",", "").replace(" ", "")
    return text


def save_plan(
    view: ModelView, owner: str | None, table: Sequence[Sequence[str]]
) -> str:
    """Keep a checked file until the import is confirmed, and name it."""
    PLAN_FOLDER.mkdir(parents=True, exist_ok=True)
    _forget_old_plans()
    token = secrets.token_urlsafe(24)
    payload = {"view": view.name, "owner": owner, "table": [list(row) for row in table]}
    (PLAN_FOLDER / f"{token}.json").write_text(json.dumps(payload), encoding="utf-8")
    return token


def load_plan(token: str, view: ModelView, owner: str | None) -> list[list[str]] | None:
    """Take back a kept file, if it is still there and belongs to this person."""
    if not token.replace("-", "").replace("_", "").isalnum():
        return None
    path = PLAN_FOLDER / f"{token}.json"
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if payload.get("view") != view.name or payload.get("owner") != owner:
        return None
    path.unlink(missing_ok=True)
    table: list[list[str]] = payload["table"]
    return table


def _forget_old_plans() -> None:
    cutoff = time.time() - PLAN_LIFETIME
    for path in PLAN_FOLDER.glob("*.json"):
        try:
            if path.stat().st_mtime < cutoff:
                path.unlink(missing_ok=True)
        except OSError:
            continue


def template_csv(view: ModelView, request: Any = None) -> str:
    """A header row naming every column a file can fill."""
    buffer = io.StringIO()
    csv.writer(buffer).writerow(import_columns(view, request))
    return buffer.getvalue()
